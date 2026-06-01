from __future__ import annotations

import io
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"C:\Users\adilb\OneDrive\Desktop\Diplom")
TEMPLATE = Path(r"D:\КСП\Нұржанұлы Саяхат - дипломдық жұмыс.docx")
SOURCE_CONTENT = ROOT / "adilbek_Diplom_jymis_keneytilgen_60pages.docx"
OUTPUT = ROOT / "outputs" / "ORDA_diplom_word_ulgisi_boyynsha.docx"

FONT = "Times New Roman"

TITLE_MAP = {
    "1 ІС-ШАРАЛАРДЫ БАСҚАРУ ВЕБ-ЖҮЙЕЛЕРІНІҢ ТЕОРИЯЛЫҚ НЕГІЗДЕРІ": "1 ІС-ШАРАЛАР ВЕБ-ЖҮЙЕСІНІҢ НЕГІЗДЕРІ",
    "Цифрландыру жағдайындағы іс-шараларды ұйымдастыру ерекшеліктері": "1.1 Цифрландыру және іс-шараларды ұйымдастыру",
    "1.1 Цифрландыру жағдайындағы іс-шараларды ұйымдастыру ерекшеліктері": "1.1 Цифрландыру және іс-шараларды ұйымдастыру",
    "1.2 Қатысушыларды тіркеу, залдарды брондау және пайдаланушы тәжірибесі": "1.2 Қатысушы тіркеуі және зал брондау",
    "1.3 Ұқсас платформаларға шолу және ORDA жүйесінің орны": "1.3 Ұқсас платформалар және ORDA орны",
    "1.4 Технологиялық стек және қауіпсіздік негіздемесі": "1.4 Технологиялық стек және қауіпсіздік",
    "2.1 Жүйеге қойылатын талаптар және пайдаланушы сценарийлері": "2.1 Талаптар және пайдаланушы сценарийлері",
    "2.3 Деректер қоры құрылымы, API логикасы және рөлдік қолжетімділік": "2.3 Деректер қоры және API логикасы",
    "2.4 Интерфейсті жобалау, көптілділік және визуалдық стиль": "2.4 Интерфейс, көптілділік және стиль",
    "3 ЖҮЙЕНІ ІСКЕ АСЫРУ, ТЕКСЕРУ ЖӘНЕ ТИІМДІЛІГІН БАҒАЛАУ": "3 ЖҮЙЕНІ ІСКЕ АСЫРУ ЖӘНЕ ТЕСТІЛЕУ",
    "3.1 Клиенттік бөліктің негізгі модульдерін жүзеге асыру": "3.1 Клиенттік модульдерді іске асыру",
    "3.2 Серверлік бөлікті, брондау логикасын және әкімші модулін іске асыру": "3.2 Серверлік бөлік және әкімші модулі",
    "3.3 Тестілеу нәтижелері және анықталған тәуекелдерді талдау": "3.3 Тестілеу және тәуекелдер",
    "3.4 Жобаның практикалық және экономикалық тиімділігі": "3.4 Практикалық және экономикалық тиімділік",
}

EXTENDED_SECTIONS = {
    "1.1": [
        "Іс-шараларды цифрландыру тек электрондық тіркеу формасын енгізумен шектелмейді. Бұл процесс қатысушылардың деректерін жинау, зал ресурстарын бөлу, уақыт кестесін бақылау, ұйымдастырушылардың шешім қабылдауын жеңілдету және барлық әрекетті бір ақпараттық кеңістікке біріктіруді қамтиды. Сондықтан веб-жүйе нақты ұйымдағы жұмыс тәртібін модельдейтін басқару құралы ретінде қарастырылады.",
        "Дәстүрлі тәсілде іс-шараға қатысушылар тізімі кесте файлдарында, зал кестесі бөлек журналда, ал өтінім мәртебесі жеке хабарламаларда сақталуы мүмкін. Мұндай ортада ақпараттың қайталануы, бір өтінімнің екі рет өңделуі немесе бос емес залдың қайта брондалуы сияқты тәуекелдер пайда болады. Веб-жүйе бұл тәуекелдерді деректерді бір базада сақтау арқылы азайтады.",
        "Цифрлық жүйенің артықшылығы пайдаланушы әрекеттерінің ізін сақтауында. Әрбір тіркеу, брондау, мәртебе өзгерісі және техникалық қолдау хабарламасы уақыт белгісімен деректер қорына жазылады. Бұл ұйымдастырушыға процестің қай кезеңде тұрғанын көруге, ал пайдаланушыға өз өтінімінің нәтижесін жеке кабинеттен бақылауға мүмкіндік береді.",
        "Іс-шараларды ұйымдастыруда жылдамдық пен сенімділік қатар маңызды. Қатысушы іс-шара туралы ақпаратты бірден көріп, бос орын бар-жоғын анықтап, тіркелуі қажет. Әкімші бұл өтінімді бөлек файлдан іздемей, бір панельде көріп, қажет болса мәртебесін өзгертеді. Осылайша цифрландыру қызмет көрсету уақытын қысқартады.",
        "ORDA Smart Event System жобасы осы теориялық талаптарды практикалық деңгейде іске асыруға бағытталған. Жүйеде каталог, күнтізбе, конференц-залдар тізімі, өтінім мәртебесі, техникалық қолдау және әкімшілік бақылау бір интерфейске біріктірілген. Мұндай тәсіл дипломдық жобаның қолданбалы маңызын арттырады.",
    ],
    "1.2": [
        "Қатысушыны тіркеу сценарийі жүйенің негізгі бизнес-процесінің бірі болып табылады. Пайдаланушы алдымен іс-шаралар каталогынан қажетті жазбаны таңдайды, содан кейін detail бетінде күнін, уақытын, форматын, мекенжайын, бағасын және бос орын санын көреді. Егер іс-шара қолжетімді болса, тіркелу батырмасы арқылы өтінім жасалады.",
        "Залды брондау сценарийі қатысушы тіркеуінен күрделірек. Мұнда тек пайдаланушы деректері ғана емес, залдың сыйымдылығы, жабдықтары, бағасы, күн мен уақыт аралығы, қосымша қызметтер және өтінім мақсаты ескеріледі. Сервер жаңа өтінім қабылдамас бұрын уақыт аралықтарының қабаттасуын тексереді.",
        "Пайдаланушы тәжірибесі қарапайым және түсінікті болуы керек. Егер интерфейс күрделі болса, пайдаланушы қажетті іс-шараны табуға немесе зал брондауға көп уақыт жұмсайды. Сондықтан іздеу өрісі, категория сүзгісі, формат таңдауы, календарь және жеке кабинет элементтері бір-бірімен логикалық байланыста орналасады.",
        "Жүйеде өтінім мәртебесінің көрінуі ерекше маңызды. Қолмен ұйымдастырылатын процесте пайдаланушы өтінімі қабылданды ма, қаралуда ма, әлде қабылданбады ма деген ақпаратты бөлек сұрауы мүмкін. ORDA жүйесінде бұл ақпарат жеке кабинет пен хабарлама модулі арқылы беріледі.",
        "Пайдаланушыға көрсетілетін әрбір экран нақты міндетті орындауға бағытталған: каталог таңдау үшін, detail бет шешім қабылдау үшін, форма өтінім жіберу үшін, жеке кабинет бақылау үшін, ал хабарлама аймағы кері байланыс үшін қолданылады. Осы құрылым веб-жүйені түсінікті және кәсіби етеді.",
    ],
    "1.3": [
        "Ұқсас платформаларды қарастырғанда олардың көпшілігі іс-шара жариялау, билет сату немесе қатысушы тіркеу функциясына басымдық беретіні байқалады. Алайда конференц-залдарды брондау, техникалық қолдау, admin бақылауы және көптілді интерфейсті бір дипломдық жобада біріктіру сирек кездеседі. ORDA жүйесінің ерекшелігі осы модульдердің бір жұмыс процесінде байланысуында.",
        "Event management жүйелерінде пайдаланушы көбіне іс-шараны көріп, тіркелумен ғана шектеледі. Ал ұйым ішінде зал ресурсын басқару бөлек құрал арқылы жүргізіледі. Бұл екі процестің бөлінуі әкімші үшін артық жұмыс тудырады. ORDA жобасы іс-шара мен зал брондау логикасын бір Booking моделінің түрлері арқылы біріктіреді.",
        "Конференц-орталықтарда залдың бос уақытын тексеру маңызды. Егер жүйе уақыт аралықтарын автоматты тексермесе, бір залға бірнеше өтінім түсуі мүмкін. ORDA жобасында availability логикасы осы мәселені шешуге бағытталған және нақты ұйымда қолдануға болатын практикалық шешім ұсынады.",
        "Көптілділік ұқсас жүйелермен салыстырғанда маңызды артықшылық болып табылады. Қазақстан жағдайында қазақ, орыс және ағылшын тілдерін қолдау пайдаланушылардың әртүрлі топтарына ыңғайлы орта қалыптастырады. Интерфейс мәтіндерінің i18n құрылымында сақталуы жүйені кейін кеңейтуді жеңілдетеді.",
        "Осы салыстырмалы талдау ORDA Smart Event System жобасының оқу мақсатынан бөлек нақты қолдануға жақын екенін көрсетеді. Жүйе іс-шаралар, залдар, өтінімдер, пайдаланушылар және қолдау хабарламалары арасындағы байланысты бір архитектураға жинақтайды.",
    ],
    "1.4": [
        "Технологиялық стек таңдау дипломдық жобаның тұрақтылығы мен түсініктілігіне әсер етеді. HTML, CSS және JavaScript frontend бөлігін браузерде қосымша құрастыру құралдарынсыз іске қосуға мүмкіндік береді. Бұл тәсіл жобаны қорғау кезінде демонстрацияны жеңілдетеді және интерфейстің негізгі логикасын ашық көрсетеді.",
        "Node.js және Express.js серверлік бөлік үшін таңдалды, себебі олар REST API құруға, middleware қабатын ұйымдастыруға және статикалық frontend файлдарын таратуға ыңғайлы. Express маршрутизациясы әр модульді жеке route файлына бөлуге мүмкіндік береді, бұл код құрылымын түсінікті етеді.",
        "MongoDB Atlas бұлттық деректер қоры ретінде қолданылды. Құжаттық модель іс-шара, зал, пайдаланушы және брондау сияқты объектілерді икемді сипаттауға қолайлы. Mongoose схемалары өрістердің түрін, міндеттілігін, enum мәндерін және бастапқы мәндерін анықтайды.",
        "Қауіпсіздік негізінде JWT авторизациясы жатыр. Пайдаланушы жүйеге кірген кезде сервер қол қойылған token қайтарады, ал қорғалған маршруттарда бұл token тексеріледі. Әкімші функциялары adminOnly middleware арқылы бөлініп, рөлдік қолжетімділік сақталады.",
        "Технологиялық стек жобаның болашақта кеңеюіне де мүмкіндік береді. REST API сыртқы мобильді қосымшаға немесе жеке admin интерфейске қызмет көрсете алады, MongoDB Atlas деректер көлемі артқанда масштабталуы мүмкін, ал frontend модульдері жаңа беттермен толықтырылады.",
    ],
    "2.1": [
        "Талаптарды анықтау кезеңінде жүйенің кім үшін жасалатыны нақты белгіленеді. ORDA платформасында негізгі пайдаланушыларға іс-шараға қатысушы, зал брондаушы клиент және әкімші жатады. Әр рөлдің интерфейстегі мүмкіндігі мен серверлік қолжетімділігі әртүрлі.",
        "Қарапайым пайдаланушы үшін басты талаптар - тіркелу, жүйеге кіру, каталогты қарау, іс-шараға жазылу, залға өтінім беру және жеке кабинеттен мәртебені бақылау. Бұл сценарийлер жүйенің клиенттік бөлігі мен API endpoint-тері арасында тұрақты байланысты қажет етеді.",
        "Әкімші үшін талаптар кеңірек: іс-шара қосу, өңдеу, жою, залдарды басқару, өтінім мәртебесін өзгерту, check-in белгілеу, статистика қарау және support хабарламаларын оқу. Бұл әрекеттер рөлдік қолжетімділікпен шектелуі тиіс, себебі олар жүйедегі негізгі деректерге әсер етеді.",
        "Функционалдық емес талаптар да маңызды. Жүйе әртүрлі экран өлшемінде дұрыс көрінуі, сұраныстарға жылдам жауап беруі, қате жағдайларын түсінікті хабарламамен көрсетуі және деректерді жоғалтпай сақтауы керек. Мұндай талаптар пайдаланушы сенімін қалыптастырады.",
        "Талаптар жүйені тек экран жиынтығы ретінде емес, толық ақпараттық процесс ретінде қарастыруға мүмкіндік береді. Әр талап нақты модульмен байланысады: авторизация auth controller арқылы, каталог event controller арқылы, залдар hall controller арқылы, ал өтінімдер booking controller арқылы орындалады.",
    ],
    "2.2": [
        "Клиент-сервер архитектурасында frontend пайдаланушы әрекетін қабылдайды, ал backend деректерді тексеріп, бизнес-логиканы орындайды. Бұл бөлініс интерфейстің көрінісін серверлік есептеулерден ажыратады және жүйені кезең-кезеңімен дамытуға мүмкіндік береді.",
        "Frontend HTTP сұраныстарды арнайы helper арқылы жібереді. Бұл тәсіл әр жерде fetch логикасын қайталамай, авторизация token-ін, қате өңдеуді және JSON жауаптарын бір тәртіпке келтіреді. Нәтижесінде код құрылымы ықшам әрі қолдауға ыңғайлы болады.",
        "Backend Express қосымшасы бірнеше қабаттан тұрады: route сұранысты қабылдайды, middleware қауіпсіздік пен валидацияны тексереді, controller бизнес-логиканы орындайды, ал model деректер қорымен жұмыс істейді. Мұндай құрылым MVC қағидасына жақын және дипломдық жобада түсіндіруге ыңғайлы.",
        "Деректер ағыны нақты мысал арқылы жақсы көрінеді. Пайдаланушы зал брондау формасын толтырғанда frontend `/api/hall-bookings` endpoint-іне сұраныс жібереді. Сервер пайдаланушы token-ін тексереді, уақыт аралықтарының қабаттасуын анықтайды, жаңа Booking құжатын сақтайды және JSON жауап қайтарады.",
        "Архитектурадағы маңызды шешімдердің бірі - frontend пен backend бір жобада орналасса да, API namespace арқылы байланысуы. Бұл шешім кейін production деңгейінде қабаттарды бөлек орналастыруға мүмкіндік береді және қазіргі кодты толық қайта жазуды қажет етпейді.",
    ],
    "2.3": [
        "Деректер қоры құрылымы жүйенің тұрақты жұмысын анықтайды. Егер модельдер дұрыс бөлінбесе, бір объектіге қатысты ақпарат бірнеше жерде қайталанып, кейін жаңарту қиындайды. ORDA жүйесінде пайдаланушы, іс-шара, зал, брондау және қолдау хабарламасы жеке модельдер ретінде қарастырылған.",
        "User моделінде пайдаланушының аты, email мекенжайы, құпиясөзі, рөлі, телефоны, ұйымы және қаласы сақталады. Email бірегей болуы тиіс, себебі авторизация осы өріс арқылы орындалады. Role өрісі user және admin мәндері арқылы қолжетімділікті басқаруға негіз болады.",
        "Event моделінде іс-шараның атауы, күні, аяқталу күні, уақыты, категориясы, сипаттамасы, ұйымдастырушысы, орны, форматы, бағасы, сыйымдылығы және жариялану мәртебесі бар. Бұл өрістер каталог карточкасын, detail бетін және admin формасын қалыптастыруға жеткілікті.",
        "Hall моделінде зал атауы, қабаты, орналасқан жері, сыйымдылығы, сағаттық бағасы, жабдықтары, сипаттамасы, аудармалары, мәртебесі және артықшылықтары сақталады. Бұл құрылым конференц-залды тек атау ретінде емес, брондауға болатын ресурс ретінде сипаттайды.",
        "Booking моделі жүйедегі ең күрделі объектілердің бірі. Ол event, hall және custom-event түрлерін қолдайды, пайдаланушы деректерін, уақыт аралығын, төлем мәртебесін, admin ескертпесін және check-in белгісін сақтайды. Осы модель арқылы қатысушы тіркеуі мен зал брондау бір логикада өңделеді.",
    ],
    "2.4": [
        "Интерфейсті жобалауда негізгі мақсат - пайдаланушыға әрекетті тез орындауға мүмкіндік беру. Басты бет жүйенің жалпы мәнін көрсетеді, каталог қажетті іс-шараны табуға көмектеседі, залдар бөлімі ресурстарды салыстыруға мүмкіндік береді, ал жеке кабинет өтінімдерді бақылауға арналған.",
        "Көптілділік интерфейстің кәсіби деңгейін арттырады. Қазақ, орыс және ағылшын тілдерінің болуы оқу орны, конференц-орталық немесе бизнес алаңы сияқты әртүрлі ұйымдарға жүйені бейімдеуге жағдай жасайды. Мәтіндер i18n сөздігінде сақталып, таңдалған тілге қарай динамикалық ауысады.",
        "Визуалдық стильде ақпараттың оқылуына басымдық беріледі. Батырмалар, карточкалар, форма өрістері және кестелер бір жүйелі дизайнға бағынады. Бұл пайдаланушыға әр бөлімде жаңа логиканы қайта үйренбей, бұрынғы әрекет үлгісін қолдануға мүмкіндік береді.",
        "Жеке кабинет интерфейсі пайдаланушы үшін ерекше маңызды, себебі ол жүйемен өзара әрекетінің нәтижесін осы жерден көреді. Мұнда тіркелген іс-шаралар, зал өтінімдері, мәртебе және хабарламалар көрсетіледі. Осындай кері байланыс жүйенің сенімділігін арттырады.",
        "Admin интерфейсте ақпарат тығыз әрі басқаруға ыңғайлы орналасуы керек. Әкімші бірнеше өтінімді салыстырады, мәртебе өзгертеді, статистика қарайды және support хабарламаларына жауап береді. Сондықтан бұл бөлімде артық визуалды элементтерден гөрі нақты деректердің көрінуі маңызды.",
    ],
    "3.1": [
        "Клиенттік бөлікті іске асыруда басты файлдар ретінде `index.html`, `styles.css` және `app.js` қолданылды. HTML құжат беттердің құрылымын береді, CSS интерфейстің визуалдық көрінісін қалыптастырады, ал JavaScript деректерді жүктеу, формаларды өңдеу және бөлімдер арасында ауысу логикасын орындайды.",
        "Каталог модулі event деректерін серверден алып, оларды карточка түрінде көрсетеді. Әр карточкада іс-шара атауы, күні, форматы, категориясы, орны және тіркелу мүмкіндігі беріледі. Іздеу және сүзгілеу элементтері пайдаланушыға көп жазба арасынан қажеттісын тез табуға көмектеседі.",
        "Конференц-залдар модулі Hall деректеріне сүйенеді. Зал карточкасында сыйымдылық, қабат, орналасқан жер, жабдықтар және баға көрсетіледі. Пайдаланушы нақты залды таңдап, брондау формасын аша алады. Бұл сценарий ресурс басқару міндетін клиенттік деңгейде түсінікті етеді.",
        "Календарь модулі іс-шараларды күн бойынша көруге мүмкіндік береді. Мұндай көрініс ұйымдастырушы үшін де, қатысушы үшін де пайдалы, себебі бір күндегі бірнеше іс-шараны салыстыру жеңілдейді. Календарь іздеу мен сүзгілеу құралдарымен бірге жұмыс істейді.",
        "Frontend формалары пайдаланушы енгізген деректерді бастапқы деңгейде тексереді, бірақ негізгі тексеру серверде орындалады. Бұл қауіпсіздік үшін қажет, себебі браузердегі тексеруді айналып өту мүмкін. Сондықтан frontend ыңғайлылықты, ал backend деректер тұтастығын қамтамасыз етеді.",
    ],
    "3.2": [
        "Серверлік бөлік Express қосымшасы ретінде ұйымдастырылды. Негізгі app файлы CORS, JSON parser, static файлдарды тарату және API route-тарын қосады. Бұл қабат барлық сұраныстардың бір кіру нүктесі болып табылады және жүйенің backend тәртібін анықтайды.",
        "Auth controller пайдаланушыны тіркеу және жүйеге кіргізу міндетін атқарады. Тіркелу кезінде email бірегейлігі тексеріледі, құпиясөз хэштеледі, ал сәтті кіргеннен кейін JWT token беріледі. Бұл token кейін қорғалған сұраныстарда қолданылады.",
        "Booking controller іс-шараға тіркелу, зал брондау, төлем мәртебесін жаңарту, өтінімді жою, хабарламаларды оқу және availability тексеру логикасын қамтиды. Бұл controller жүйедегі ең негізгі бизнес-процестерді біріктіреді.",
        "Admin controller әкімші әрекеттерін басқарады. Ол барлық өтінімдерді көруге, пайдаланушылар тізімін алуға, статистика есептеуге, check-in белгілеуге, өтінім мәртебесін өзгертуге және demo деректерді жаңартуға мүмкіндік береді. Бұл модуль ұйым қызметкерінің күнделікті жұмысын автоматтандырады.",
        "Support controller пайдаланушы хабарламаларын қабылдап, admin панельде көрсету үшін қолданылады. Техникалық қолдау модулі пайдаланушы мен ұйымдастырушы арасындағы қосымша байланыс арнасын қалыптастырады және жүйенің қызмет көрсету сапасын арттырады.",
    ],
    "3.3": [
        "Тестілеу жүйенің тек сәтті сценарийлерін емес, қате жағдайларын да қамтуы тиіс. Мысалы, бос email арқылы тіркелу, жарамсыз token жіберу, admin маршрутын қарапайым пайдаланушымен ашу, бос емес залға өтінім беру сияқты жағдайлар тексеріледі.",
        "Frontend деңгейінде синтаксистік қателерді анықтау үшін `npm test` командасы қолданылды. Бұл команда негізгі JavaScript файлдарының дұрыс жазылғанын тексереді. Мұндай тексеру браузер ашпай тұрып-ақ бастапқы кодтағы қателерді табуға мүмкіндік береді.",
        "Backend деңгейінде endpoint жауаптары, status code мәндері, JSON құрылымы және middleware жұмысы қарастырылады. Auth middleware token жоқ кезде 401 қатесін, adminOnly middleware құқық жеткіліксіз болғанда 403 қатесін қайтаруы керек.",
        "Деректер қорымен байланысты тәуекелдер де маңызды. MongoDB Atlas қолжетімсіз болса, events, halls және bookings деректері жүктелмейді. Сондықтан health endpoint арқылы database күйін тексеру және `.env` параметрлерін дұрыс сақтау қажет.",
        "Тестілеу нәтижелері жүйенің негізгі дипломдық мақсаттарды орындайтынын көрсетті. Дегенмен болашақта автоматты unit test, integration test және end-to-end test сценарийлерін қосу жүйені production деңгейіне жақындатады.",
    ],
    "3.4": [
        "Жобаның практикалық тиімділігі ұйымдастырушының жұмыс уақытын үнемдеуінен көрінеді. Өтінімдерді бір тізімнен қарау, мәртебені батырма арқылы өзгерту, check-in белгілеу және статистика алу қолмен жүргізілетін жұмыстарды азайтады.",
        "Пайдаланушы үшін тиімділік ақпараттың қолжетімді болуымен байланысты. Ол іс-шаралар каталогын қарайды, залдарды салыстырады, өтінім жібереді және нәтижесін жеке кабинеттен көреді. Бұл қызметтің ашықтығын арттырады.",
        "Экономикалық тұрғыдан жүйе қымбат лицензиялық құралдарды қажет етпейді. Ашық веб-технологиялар, Node.js, Express.js және MongoDB Atlas-тың бастапқы тегін мүмкіндіктері шағын ұйымдарға жобаны төмен шығынмен іске қосуға мүмкіндік береді.",
        "Жүйенің даму мүмкіндігі жоғары. Болашақта QR check-in, email хабарлама, SMS сервис, онлайн төлем, PDF есептер, зал жүктемесін талдау және сыртқы календарьмен синхрондау функцияларын қосуға болады.",
        "Осылайша ORDA Smart Event System дипломдық жоба ретінде ғана емес, нақты ұйымда енгізуге болатын бастапқы платформа ретінде де бағаланады. Оның архитектурасы жаңа функциялар қосуға, интерфейсі пайдаланушыға түсінікті жұмыс істеуге бағытталған.",
    ],
}


def set_run_font(run, size: float = 14, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(key), FONT)


def set_style_font(style, size: float = 14, bold: bool | None = None) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(key), FONT)


def set_cell_text(cell, text: str, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT, size: float = 14) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{name}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_table_borders(table, size: int = 8) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{name}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def set_cell_bottom_border(cell, size: int = 8) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), "000000")


def set_table_widths(table, widths_cm: list[float]) -> None:
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 567)))
            tc_w.set(qn("w:type"), "dxa")


def add_page_field(paragraph) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, sep, text, end])
    set_run_font(run, 14)


def add_toc_field(paragraph) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.left_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Мазмұн Word бағдарламасында жаңартылады"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, sep, placeholder, end])
    set_run_font(run, 14)


def set_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def configure_document(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.header_distance = Cm(1.27)
        section.footer_distance = Cm(1.27)

    normal = doc.styles["Normal"]
    set_style_font(normal, 14, False)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        set_style_font(style, 14, True)
        style.paragraph_format.line_spacing = 1
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    doc.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["Heading 1"].paragraph_format.first_line_indent = Cm(0)
    doc.styles["Heading 2"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.styles["Heading 2"].paragraph_format.first_line_indent = Cm(0)

    if "FrontMatterHeading" not in doc.styles:
        front = doc.styles.add_style("FrontMatterHeading", WD_STYLE_TYPE.PARAGRAPH)
    else:
        front = doc.styles["FrontMatterHeading"]
    front.base_style = normal
    set_style_font(front, 14, True)
    front.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    front.paragraph_format.first_line_indent = Cm(0)
    front.paragraph_format.line_spacing = 1
    front.paragraph_format.space_before = Pt(0)
    front.paragraph_format.space_after = Pt(0)
    front.paragraph_format.keep_with_next = True

    if "CaptionCustom" not in doc.styles:
        caption = doc.styles.add_style("CaptionCustom", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = doc.styles["CaptionCustom"]
    caption.base_style = normal
    set_style_font(caption, 12, False)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.line_spacing = 1
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.keep_with_next = True

    if "FigurePlaceholder" not in doc.styles:
        fig_style = doc.styles.add_style("FigurePlaceholder", WD_STYLE_TYPE.PARAGRAPH)
    else:
        fig_style = doc.styles["FigurePlaceholder"]
    fig_style.base_style = normal
    set_style_font(fig_style, 13, False)
    fig_style.font.italic = True
    fig_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig_style.paragraph_format.first_line_indent = Cm(0)
    fig_style.paragraph_format.line_spacing = 1
    fig_style.paragraph_format.space_before = Pt(6)
    fig_style.paragraph_format.space_after = Pt(3)

    for name in ("TOC 1", "TOC 2"):
        if name in doc.styles:
            style = doc.styles[name]
            set_style_font(style, 14, False)
            style.paragraph_format.line_spacing = 1
            style.paragraph_format.first_line_indent = Cm(0)
            style.paragraph_format.left_indent = Cm(0)
            style.paragraph_format.right_indent = Cm(0)
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(0)
            style.paragraph_format.tab_stops.clear_all()
            style.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


def add_centered(doc: Document, text: str, bold: bool = False, size: float = 14, before: float = 0, after: float = 0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    set_run_font(r, size, bold)
    return p


def add_body(doc: Document, text: str):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, 14)
    return p


def add_list_item(doc: Document, text: str):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("— " + text.rstrip(";.") + ";")
    set_run_font(r, 14)
    return p


def add_reference(doc: Document, text: str):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, 14)
    return p


def add_code_title(doc: Document, text: str):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, 14, True)
    return p


def add_code_line(doc: Document, text: str):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, 10.5)
    return p


def wrap_code_line(line: str, width: int = 102) -> list[str]:
    line = line.replace("\t", "  ").rstrip()
    if not line:
        return [""]
    chunks: list[str] = []
    prefix = ""
    while len(line) > width:
        split_at = max(line.rfind(" ", 0, width), line.rfind(",", 0, width), line.rfind(")", 0, width))
        if split_at < 50:
            split_at = width
        chunks.append(prefix + line[:split_at].rstrip())
        line = "    " + line[split_at:].lstrip()
        prefix = ""
    chunks.append(prefix + line)
    return chunks


def add_code_file(doc: Document, title: str, relative_path: str, max_lines: int | None = None) -> None:
    path = ROOT / relative_path
    add_code_title(doc, title)
    if not path.exists():
        add_code_line(doc, f"Файл табылмады: {relative_path}")
        return
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    if max_lines is not None:
        lines = lines[:max_lines]
    add_code_line(doc, f"// Файл: {relative_path}")
    for index, line in enumerate(lines, start=1):
        wrapped = wrap_code_line(line)
        for part_index, part in enumerate(wrapped):
            label = f"{index:04d}  " if part_index == 0 else "      "
            add_code_line(doc, label + part)


def add_heading1(doc: Document, text: str, page_break: bool = True):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(text)
    set_run_font(r, 14, True)
    return p


def add_heading2(doc: Document, text: str):
    p = doc.add_paragraph(style="Heading 2")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r, 14, True)
    return p


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, 6)
    set_table_widths(table, widths)
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_text(cell, text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    for row_data in rows:
        row = table.add_row()
        for cell, text in zip(row.cells, row_data):
            set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=12)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    return table


def add_screenshot_placeholder(doc: Document, number: int, caption: str, reference: str | None = None, height_cm: float = 7.6):
    if reference:
        add_body(doc, reference)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, 6)
    set_table_widths(table, [15.5])
    row = table.rows[0]
    row.height = Cm(height_cm)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("Скриншот интерфейса")
    set_run_font(run, 13, False, True)
    cap = doc.add_paragraph(f"Сурет {number}. {caption}", style="CaptionCustom")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


def extract_logo_bytes() -> io.BytesIO | None:
    with zipfile.ZipFile(TEMPLATE) as zf:
        for name in zf.namelist():
            if name == "word/media/image1.jpeg":
                return io.BytesIO(zf.read(name))
    return None


def add_cover(doc: Document) -> None:
    logo = extract_logo_bytes()
    header = doc.add_table(rows=1, cols=2)
    remove_table_borders(header)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(header, [3.2, 13.4])
    left, right = header.rows[0].cells
    if logo:
        p_logo = left.paragraphs[0]
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.first_line_indent = Cm(0)
        run = p_logo.add_run()
        run.add_picture(logo, width=Cm(2.2), height=Cm(2.2))
    text = (
        "ҚР ОҚУ-АҒАРТУ МИНИСТРЛІГІ\n"
        "МИНИСТЕРСТВО ПРОСВЕЩЕНИЯ РК\n\n"
        "АЛМАТЫ ҚАРЖЫ-ЭКОНОМИКАЛЫҚ КОЛЛЕДЖІ\n"
        "АЛМАТИНСКИЙ ФИНАНСОВО-ЭКОНОМИЧЕСКИЙ\n"
        "КОЛЛЕДЖ"
    )
    right.text = ""
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    for i, line in enumerate(text.split("\n")):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, 13.5, True)

    line = doc.add_paragraph()
    line.paragraph_format.first_line_indent = Cm(0)
    line.paragraph_format.space_before = Pt(8)
    line.paragraph_format.space_after = Pt(24)
    p_pr = line._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "double")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    p_pr.append(borders)

    for text in (
        "«Қорғауға рұқсат»",
        "Колледж директоры",
        "__________ Алимова К.У.",
        "«____» _____________2026 ж.",
        "хаттама №____",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(10.4)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1
        r = p.add_run(text)
        set_run_font(r, 14)

    add_centered(doc, "ДИПЛОМДЫҚ ЖҰМЫС", size=14, before=24, after=10)
    add_centered(doc, "Тақырыбы «Іс-шаралар мен конференцияларды тіркеу және басқаруға арналған веб-жүйе»", size=14)

    meta = doc.add_table(rows=2, cols=2)
    remove_table_borders(meta)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(meta, [5.2, 10.4])
    rows = [
        ("Мамандық", "06130100 Бағдарламалық қамтамасыз ету\n(түрлері бойынша)"),
        ("Біліктілік", "4S06130105 Ақпараттық жүйелер технигі"),
    ]
    for row, (label, value) in zip(meta.rows, rows):
        set_cell_text(row.cells[0], label)
        set_cell_text(row.cells[1], value)
        set_cell_bottom_border(row.cells[1])

    p_gap = doc.add_paragraph()
    p_gap.paragraph_format.space_after = Pt(42)

    sig = doc.add_table(rows=2, cols=3)
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig.autofit = False
    set_table_borders(sig, 6)
    set_table_widths(sig, [5.2, 5.2, 5.2])
    signature_rows = [
        ("Орындады:", "________________", "Жанұзақ Ә.Т."),
        ("Ғылыми жетекші:", "________________", "Тұрғантай А.Ә."),
    ]
    for row, values in zip(sig.rows, signature_rows):
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value, bold=False)

    add_centered(doc, "Алматы, 2026", before=118)
    doc.add_page_break()


def add_front_matter(doc: Document) -> None:
    annotations = [
        (
            "АҢДАТПА",
            "Бұл дипломдық жұмыста іс-шараларға қатысушыларды тіркеу, конференц-залдарды брондау, өтінімдерді өңдеу және әкімшілік бақылауды автоматтандыруға арналған веб-жүйені жобалау мен іске асыру мәселелері қарастырылды. Жоба HTML, CSS, JavaScript, Node.js, Express.js, MongoDB Atlas, REST API және JWT авторизациясы негізінде әзірленді. Жүйеде көптілді интерфейс, пайдаланушыны тіркеу, іс-шаралар каталогы, залдарды брондау, күнтізбе, хабарламалар, техникалық қолдау және әкімшілік панель модульдері қамтылды.",
        ),
        (
            "АННОТАЦИЯ",
            "В дипломной работе рассмотрены проектирование и реализация веб-системы для регистрации участников мероприятий, управления конференциями и бронирования конференц-залов. Система разработана на основе HTML, CSS, JavaScript, Node.js, Express.js, MongoDB Atlas, REST API и JWT-авторизации. В работе раскрыты архитектура клиент-серверного взаимодействия, структура базы данных, многоязычный интерфейс, каталог мероприятий, календарь, система заявок, уведомления, техническая поддержка и административная панель.",
        ),
        (
            "ABSTRACT",
            "This diploma thesis describes the design and implementation of a web system for event registration, conference management and conference hall booking. The project is built with HTML, CSS, JavaScript, Node.js, Express.js, MongoDB Atlas, REST API and JWT authorization. The thesis covers client-server architecture, database models, multilingual interface, event catalog, booking workflow, calendar, request management, notifications, technical support and administrative dashboard.",
        ),
    ]
    for title, text in annotations:
        p = doc.add_paragraph(style="FrontMatterHeading")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        set_run_font(r, 14, True)
        add_body(doc, text)
    doc.add_page_break()

    title = doc.add_paragraph(style="FrontMatterHeading")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    r = title.add_run("МАЗМҰНЫ")
    set_run_font(r, 14, False)
    toc = doc.add_paragraph()
    add_toc_field(toc)
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.footer.is_linked_to_previous = False
    add_page_field(section.footer.paragraphs[0])


def collect_source_items() -> list[tuple[str, str]]:
    source = Document(SOURCE_CONTENT)
    items: list[tuple[str, str]] = []
    in_body = False
    for p in source.paragraphs:
        text = " ".join(p.text.split())
        if not text:
            continue
        lowered = text.lower()
        if lowered.endswith("-сурет.") or "-сурет." in lowered or "-кесте." in lowered:
            continue
        style = p.style.name if p.style else "Normal"
        if text == "КІРІСПЕ":
            in_body = True
        if not in_body:
            continue
        if text == "ҚОРЫТЫНДЫ":
            break
        if style.startswith("Heading 1"):
            items.append(("h1", TITLE_MAP.get(text, text)))
        elif style.startswith("Heading 2"):
            if text == "Цифрландыру жағдайындағы іс-шараларды ұйымдастыру ерекшеліктері":
                text = "1.1 Цифрландыру жағдайындағы іс-шараларды ұйымдастыру ерекшеліктері"
            items.append(("h2", TITLE_MAP.get(text, text)))
        elif style.startswith("List"):
            items.append(("list", text))
        else:
            items.append(("p", text))
    return items


def add_extended_for_heading(doc: Document, heading: str) -> None:
    key = heading.split()[0] if heading else ""
    for text in EXTENDED_SECTIONS.get(key, []):
        add_body(doc, text)


def maybe_insert_hook(doc: Document, heading: str) -> None:
    if heading.startswith("1.2"):
        add_screenshot_placeholder(
            doc,
            1,
            "Іс-шара detail бетінің интерфейсі",
            "Пайдаланушы іс-шара туралы толық мәліметті detail бетінен көреді. Іс-шара detail беті 1-суретте көрсетілген.",
        )
    elif heading.startswith("2.1"):
        add_screenshot_placeholder(
            doc,
            2,
            "Пайдаланушыны тіркеу және авторизация терезесі",
            "Жүйеде пайдаланушының бастапқы әрекеті тіркеу немесе кіру формасынан басталады. Авторизация терезесі 2-суретте көрсетілген.",
        )
        add_table(
            doc,
            ["Талап түрі", "Мазмұны", "Жүйедегі іске асуы"],
            [
                ["Функционалдық", "Тіркеу, авторизация, каталог, брондау, өтінім мәртебесі", "Frontend формалары және REST API endpoint-тері"],
                ["Қауіпсіздік", "Жеке деректерді қорғау және рөлдік қолжетімділік", "JWT token, adminOnly middleware, парольді хэштелген түрде сақтау"],
                ["Қолжетімділік", "Қазақ, орыс және ағылшын тіліндегі интерфейс", "i18n сөздіктері және интерфейс мәтіндерін динамикалық ауыстыру"],
                ["Кеңейтілу", "Жаңа іс-шара, зал, хабарлама және есеп модульдерін қосу", "Express routes, Mongoose модельдері және модульдік frontend құрылымы"],
            ],
            [3.4, 5.8, 6.3],
        )
    elif heading.startswith("2.2"):
        add_screenshot_placeholder(
            doc,
            3,
            "Жүйенің басты беті және жұмыс кеңістігі",
            "Клиенттік және серверлік қабаттардың нәтижесі пайдаланушыға басты бет арқылы ұсынылады. Жүйенің басты беті 3-суретте көрсетілген.",
        )
    elif heading.startswith("2.3"):
        add_table(
            doc,
            ["Модель", "Негізгі өрістер", "Мақсаты"],
            [
                ["User", "name, email, password, role, phone, organization, city", "Пайдаланушы профилін және рөлін сақтау"],
                ["Event", "title, date, category, format, location, capacity, status", "Іс-шаралар каталогын қалыптастыру"],
                ["Hall", "name, floor, location, capacity, pricePerHour, equipment, status", "Конференц-залдар туралы деректерді сақтау"],
                ["Booking", "type, userEmail, eventId, hallId, date, status, paymentStatus", "Іс-шараға тіркелу және зал брондау өтінімдерін басқару"],
                ["SupportMessage", "userName, userEmail, text, status, readAt", "Техникалық қолдау хабарламаларын өңдеу"],
            ],
            [2.7, 6.3, 6.5],
        )
        add_screenshot_placeholder(
            doc,
            4,
            "Жеке кабинеттегі өтінімдер тізімі",
            "Деректер қоры мен API жұмысының нәтижесі жеке кабинетте нақты жазбалар түрінде көрінеді. Пайдаланушының өтінімдері 4-суретте көрсетілген.",
        )
    elif heading.startswith("2.4"):
        add_screenshot_placeholder(
            doc,
            5,
            "Көптілді интерфейс және негізгі навигация",
            "Интерфейс мәтіндері таңдалған тілге байланысты ауысып отырады. Көптілді навигация 5-суретте көрсетілген.",
        )
    elif heading.startswith("3.1"):
        add_screenshot_placeholder(
            doc,
            6,
            "Іс-шаралар каталогы және сүзгілеу құралдары",
            "Каталогта іздеу, категория және формат бойынша сүзгілеу қарастырылған. Іс-шаралар каталогы 6-суретте көрсетілген.",
        )
        add_body(doc, "Іс-шаралар каталогы пайдаланушыға жүйедегі негізгі ақпаратты бірден көруге мүмкіндік береді. Карточка құрылымында атау, күн, формат, орын, баға және бос орын саны беріледі, сондықтан пайдаланушы шешімді бірнеше қосымша бет ашпай-ақ қабылдай алады.")
        add_body(doc, "Каталогтағы сүзгілеу құралдары деректер көлемі көбейген кезде ерекше маңызды болады. Search, category және format параметрлері frontend арқылы API сұранысына жіберіліп, серверден нақты шартқа сәйкес келетін events тізімі алынады.")
        doc.add_page_break()
        add_screenshot_placeholder(
            doc,
            7,
            "Конференц-залды брондау формасы",
            "Залға өтінім беру кезінде пайдаланушы күнді, уақытты, қатысушылар санын және мақсатты көрсетеді. Брондау формасы 7-суретте көрсетілген.",
        )
        add_body(doc, "Брондау формасында енгізілетін деректер нақты ұйымдастыру процесімен байланысты. Күн, басталу уақыты, аяқталу уақыты, қатысушылар саны және мақсат өрістері кейін серверлік деңгейде тексеріліп, залдың қолжетімділігімен салыстырылады.")
        add_body(doc, "Формадағы мәліметтер толық және түсінікті болған сайын әкімші өтінімді тезірек өңдей алады. Сондықтан интерфейсте міндетті өрістер, қосымша қызметтер және төлем мәртебесі логикалық тәртіппен орналастырылған.")
        doc.add_page_break()
        add_screenshot_placeholder(
            doc,
            8,
            "Календарь және жоспарланған іс-шаралар тізімі",
            "Жоспарланған іс-шараларды уақыт бойынша бақылау үшін календарь қолданылады. Іс-шаралар календарі 8-суретте көрсетілген.",
        )
    elif heading.startswith("3.2"):
        add_screenshot_placeholder(
            doc,
            9,
            "Әкімшілік панель және өтінім мәртебелерін басқару",
            "Әкімші өтінімдерді қарап, мәртебесін өзгертіп, статистикалық көрсеткіштерді бақылайды. Әкімшілік панель 9-суретте көрсетілген.",
        )
        add_screenshot_placeholder(
            doc,
            10,
            "Техникалық қолдау хабарламаларының интерфейсі",
            "Пайдаланушы сұрағы support модулі арқылы әкімшіге жеткізіледі. Техникалық қолдау экраны 10-суретте көрсетілген.",
        )
    elif heading.startswith("3.3"):
        doc.add_page_break()
        add_table(
            doc,
            ["Тест атауы", "Күтілетін нәтиже", "Нәтиже"],
            [
                ["Пайдаланушыны тіркеу", "Жаңа аккаунт құрылып, JWT token беріледі", "Сәтті"],
                ["Іс-шараға тіркелу", "Booking жазбасы жасалып, жеке кабинетте көрсетіледі", "Сәтті"],
                ["Залды брондау", "Уақыт аралығы тексеріліп, өтінім қабылданады", "Сәтті"],
                ["Admin мәртебе өзгертуі", "Өтінім күйі жаңарып, клиент хабарлама алады", "Сәтті"],
                ["Көптілді интерфейс", "RU/KZ/EN мәтіндері дұрыс ауысады", "Сәтті"],
            ],
            [4.6, 6.2, 3.2],
        )


def add_source_body(doc: Document) -> None:
    current_heading = ""
    counts: dict[str, int] = {}
    for kind, text in collect_source_items():
        if kind == "h1":
            if current_heading:
                add_extended_for_heading(doc, current_heading)
                current_heading = ""
            add_heading1(doc, text, page_break=(text != "КІРІСПЕ"))
        elif kind == "h2":
            if current_heading:
                add_extended_for_heading(doc, current_heading)
            add_heading2(doc, text)
            current_heading = text
            counts[current_heading] = 0
        elif kind == "list":
            add_list_item(doc, text)
            if current_heading:
                counts[current_heading] += 1
        else:
            add_body(doc, text)
            if current_heading:
                counts[current_heading] += 1
                if counts[current_heading] == 2:
                    maybe_insert_hook(doc, current_heading)
    if current_heading:
        add_extended_for_heading(doc, current_heading)


def add_labor_section(doc: Document) -> None:
    add_heading1(doc, "4 ЕҢБЕКТІ ҚОРҒАУ")
    add_heading2(doc, "4.1 Компьютермен жұмыс істеу қауіпсіздігі")
    for text in [
        "Веб-жүйені әзірлеу процесі тікелей өндірістік жабдықпен емес, компьютерлік техникамен, бағдарламалық құралдармен және интернет қызметтерімен байланысты. Сондықтан еңбекті қорғау бөлімінде бағдарламашының жұмыс орнына қойылатын санитарлық-гигиеналық, эргономикалық және ақпараттық қауіпсіздік талаптары қарастырылады.",
        "Бағдарламалық өнімді жасау кезінде қызметкер ұзақ уақыт монитор алдында отырады, мәтіндік редакторлармен, браузермен, серверлік консольмен, MongoDB Atlas басқару панелімен және тестілеу құралдарымен жұмыс істейді. Мұндай еңбек түрінде көзге түсетін жүктеме, омыртқаға түсетін статикалық күш және қол буындарының қайталанатын қозғалысы негізгі тәуекелдер қатарына жатады.",
        "Жұмыс орны жеткілікті жарықтандырылуы тиіс. Табиғи жарық сол жақтан түскені дұрыс, ал жасанды жарық көздері мониторда шағылыс тудырмауы қажет. Экран жарықтығы бөлме жарығымен теңестіріліп, өте ашық немесе өте күңгірт болмауы керек. Бұл көздің шаршауын азайтады және кодпен ұзақ жұмыс істеу кезінде қателердің алдын алады.",
        "Монитор мен көз арасындағы арақашықтық шамамен 50-70 см болуы ұсынылады. Экранның жоғарғы бөлігі пайдаланушының көз деңгейінен сәл төмен орналасқанда мойын бұлшықеттеріне түсетін күш азаяды. Пернетақта мен тінтуір жұмысшының қолы табиғи қалыпта тұратындай деңгейде орналасуы керек.",
        "Әрбір 45-60 минут жұмыс істегеннен кейін қысқа үзіліс жасау, көзге арналған жаттығулар орындау және дене қалпын өзгерту ұсынылады. Бұл талаптар дипломдық жобаны әзірлеу кезінде де маңызды, себебі бағдарламалау, интерфейсті тексеру және құжаттама дайындау бірнеше сағаттық үздіксіз жұмысқа әкелуі мүмкін.",
        "Электр қауіпсіздігі тұрғысынан компьютер, маршрутизатор және қосымша құрылғылар жарамды розеткаға қосылып, сымдар зақымданбауы тиіс. Жұмыс орнында сұйықтықты электр құрылғыларына жақын қоюға болмайды. Құрылғылардың қызып кетуін болдырмау үшін желдеткіш тесіктері жабылмауы қажет.",
        "Ақпараттық қауіпсіздік те еңбекті қорғаудың заманауи бөлігі ретінде қарастырылады. Жоба `.env` файлында сақталатын MongoDB Atlas қосылу жолын және JWT құпия кілтін қолданады. Бұл деректер ашық репозиторийге шығарылмауы, бөтен адамдарға берілмеуі және жұмыс аяқталғаннан кейін қауіпсіз сақталуы тиіс.",
    ]:
        add_body(doc, text)
    add_heading2(doc, "4.2 Жұмыс орнын ұйымдастыру талаптары")
    for text in [
        "Жұмыс орнын дұрыс ұйымдастыру жүйені әзірлеушінің өнімділігіне тікелей әсер етеді. ORDA Smart Event System жобасында frontend және backend бөліктері қатар қарастырылғандықтан, әзірлеуші бір уақытта код редакторын, браузерді, терминалды, деректер қоры панелін және құжаттама файлдарын қолданады. Сондықтан экран кеңістігін, файл құрылымын және жұмыс құралдарын реттеу маңызды.",
        "Құжаттарды, кодты және тест нәтижелерін бір папка құрылымында сақтау жобаның бақылануын жеңілдетеді. Бұл жұмыста backend бөлігі `backend/src` каталогында, frontend файлдары `docs` каталогында, ал құжаттама мен қорғау материалдары бөлек файлдарда орналасқан. Мұндай құрылым жобаны әрі қарай дамытуға ыңғайлы.",
        "Жұмыс орнының үстелі таза, қажет емес қағаздар мен кабельдерден бос болуы керек. Құрылғылардың орналасуы пайдаланушының жиі қолданатын әрекеттерін жылдам орындауына мүмкіндік беруі тиіс. Мысалы, тестілеу кезінде браузер мен терминалдың қатар орналасуы сервер қатесін тез көруге көмектеседі.",
        "Бағдарламашының креслосы биіктігі реттелетін, арқалық тірегі бар болуы қажет. Аяқ еденге толық тиіп тұруы немесе арнайы тірек қолданылуы тиіс. Дене қалпы тұрақты болғанда ұзақ жұмыс кезінде шаршау азаяды және назар бағдарламалық логикаға шоғырланады.",
        "Жүйені әзірлеу кезінде психологиялық жүктеме де ескеріледі. Қате іздеу, серверлік endpoint-терді тексеру, JWT авторизациясын баптау немесе деректер қоры байланысын түзету кейде ұзақ уақыт алады. Мұндай жағдайда тапсырмаларды шағын кезеңдерге бөлу, нұсқаларды сақтау және тест нәтижелерін жазып отыру тиімді.",
        "Еңбекті қорғау талаптарын сақтау тек әзірлеушінің денсаулығын қорғап қана қоймай, бағдарламалық өнімнің сапасын арттырады. Дұрыс ұйымдастырылған ортада кодтағы қателер азаяды, интерфейсті тексеру жүйелі өтеді және дипломдық жобаның қорытынды нәтижесі тұрақты болады.",
    ]:
        add_body(doc, text)

    add_heading2(doc, "4.3 Ақпараттық қауіпсіздік және деректерді қорғау")
    for text in [
        "ORDA Smart Event System пайдаланушылардың аты-жөні, электрондық поштасы, ұйымы, қаласы, телефон нөмірі, іс-шараға қатысу өтінімі және залды брондау мәліметтері сияқты деректерді өңдейді. Сондықтан жүйеде қауіпсіздік талаптары тек техникалық функция емес, пайдаланушы сенімінің негізгі шарты ретінде қарастырылады.",
        "Авторизация JWT токені арқылы ұйымдастырылған. Пайдаланушы жүйеге кіргеннен кейін сервер қол қойылған token береді, ал қорғалған сұраныстарда бұл token `Authorization: Bearer` тақырыбы арқылы жіберіледі. Сервер token жарамды екенін тексергеннен кейін ғана жеке кабинет, өтінімдер немесе әкімші функциялары ашылады.",
        "Әкімші құқықтары бөлек middleware арқылы тексеріледі. Бұл тәсіл қарапайым пайдаланушының іс-шараны жоюына, басқа адамдардың өтінімдерін өзгертуіне немесе статистикалық деректерді көруіне жол бермейді. Рөлдік қолжетімділік ақпараттық жүйенің тұтастығын сақтайды.",
        "Құпия сөздер ашық мәтін түрінде сақталмауы керек. Жобада парольді хэштелген түрде сақтау қағидасы қарастырылған, ал `.env` файлындағы құпия кілттер репозиторийге енгізілмейді. Бұл MongoDB Atlas қосылу жолын және JWT құпиясын қорғауға мүмкіндік береді.",
        "Деректерді қорғаудың тағы бір бағыты - енгізілетін ақпаратты тексеру. Міндетті өрістердің толтырылуы, ObjectId мәндерінің дұрыстығы, өтінім мәртебелерінің рұқсат етілген тізімге сәйкес болуы серверлік деңгейде бақыланады. Мұндай тексерулер қате немесе зиянды сұраныстардан қорғайды.",
        "Жүйе болашақта HTTPS, rate limiting, аудит журналдары, резервтік көшіру және хабарлама сервистерімен толықтырылуы мүмкін. Бұл шаралар нақты ұйым жағдайында пайдаланушы деректерінің қауіпсіздігін күшейтіп, веб-жүйенің сенімділігін арттырады.",
    ]:
        add_body(doc, text)


def add_economics_section(doc: Document) -> None:
    add_heading1(doc, "5 ЭКОНОМИКАЛЫҚ БӨЛІМ")
    add_heading2(doc, "5.1 Жобаның экономикалық және әлеуметтік мәні")
    for text in [
        "Іс-шаралар мен конференцияларды басқару саласында ақпаратты қолмен өңдеу көп уақыт алады. Қатысушылар тізімін бөлек файлда жүргізу, залдың бос уақытын телефон арқылы нақтылау, өтінім мәртебесін қолмен хабарлау және статистиканы бөлек есептеу ұйымдастырушының еңбек шығынын көбейтеді. ORDA Smart Event System осы процестерді бір веб-жүйеге біріктіру арқылы уақытты үнемдейді.",
        "Жүйенің экономикалық мәні бірнеше бағытта көрінеді. Біріншіден, іс-шараға тіркелу және залды брондау өтінімдері автоматты түрде сақталады. Екіншіден, әкімші бір панель арқылы өтінімдерді қарап, мәртебесін өзгерте алады. Үшіншіден, пайдаланушы жеке кабинетінде өз өтінімдерін көріп, қосымша байланыс шығынын азайтады.",
        "Әлеуметтік тұрғыдан жүйе оқу орындары, конференц-орталықтар, бизнес алаңдары және қоғамдық ұйымдар үшін қолжетімді цифрлық құрал ұсынады. Пайдаланушы іс-шаралар туралы ақпаратты бір жерден табады, тілін таңдай алады, қажет залды брондайды және техникалық қолдау арқылы сұрақ қояды. Бұл қызмет сапасын арттырады.",
        "Жүйенің кеңейтілу мүмкіндігі де экономикалық артықшылық береді. Node.js және Express.js негізіндегі модульдік backend жаңа endpoint қосуға ыңғайлы, ал MongoDB Atlas бұлттық деректер қоры серверлік инфрақұрылымды бөлек ұстау қажеттілігін азайтады. Сондықтан алғашқы нұсқаны аз шығынмен іске қосып, кейін нақты ұйым талабына қарай кеңейтуге болады.",
        "Дайын жүйені қолдану қағаз құжат айналымын азайтады, ақпаратты қайталап енгізуді қысқартады және әкімші жұмысын жүйелейді. Бұл тікелей қаржылық үнеммен қатар, қызметкерлер уақытын тиімді бөлуге мүмкіндік береді.",
    ]:
        add_body(doc, text)

    add_heading2(doc, "5.2 Жобаны әзірлеу шығындарын есептеу")
    for text in [
        "Дипломдық жоба оқу мақсатына арналғандықтан, көптеген бағдарламалық құралдар тегін немесе ашық лицензиямен қолданылды. HTML, CSS, JavaScript, Node.js, Express.js және MongoDB драйверлері үшін бөлек лицензиялық төлем қажет емес. Негізгі шығындар әзірлеушінің уақытына, компьютер амортизациясына, интернетке және бұлттық қызметтердің ықтимал тарифтеріне байланысты есептеледі.",
        "Есептеу кезінде екі айлық әзірлеу кезеңі алынды. Бұл кезеңге талаптарды талдау, интерфейс жасау, backend логикасын жазу, MongoDB Atlas-пен байланыстыру, тестілеу, құжат дайындау және қорғауға арналған демонстрациялық сценарийді тексеру кіреді.",
    ]:
        add_body(doc, text)
    doc.add_page_break()
    add_table(
        doc,
        ["№", "Шығын атауы", "Есептеу жолы", "Сома"],
        [
            ["1", "Компьютер амортизациясы", "350 000 тг / 60 ай × 2 ай", "11 666 тг"],
            ["2", "Интернет шығыны", "7 000 тг × 2 ай", "14 000 тг"],
            ["3", "Электр энергиясы", "орташа 2 500 тг × 2 ай", "5 000 тг"],
            ["4", "Бұлттық деректер қоры", "MongoDB Atlas free tier", "0 тг"],
            ["5", "Домен және хостинг резерві", "демонстрациялық орналастыруға шартты шығын", "15 000 тг"],
            ["6", "Құжаттама және басып шығару", "дипломдық материалдарды дайындау", "8 000 тг"],
            ["", "Барлығы", "", "53 666 тг"],
        ],
        [1.1, 4.9, 6.0, 3.0],
    )
    for text in [
        "Кестеде көрсетілген шығындар жобаның оқу және демонстрациялық нұсқасына арналған. Нақты ұйымға енгізу кезінде домен, тұрақты сервер, SMS немесе email хабарлама сервисі, техникалық сүйемелдеу және резервтік көшіру шығындары бөлек есептелуі мүмкін.",
        "Жобаның бастапқы құны төмен болуы оның технологиялық стек таңдауымен байланысты. Ашық веб-технологиялар мен бұлттық деректер қоры шағын ұйымдарға жүйені үлкен инфрақұрылымсыз іске қосуға мүмкіндік береді.",
    ]:
        add_body(doc, text)

    add_heading2(doc, "5.3 Жобаның тиімділігі және даму мүмкіндіктері")
    for text in [
        "ORDA Smart Event System енгізілген жағдайда ұйымдастырушының негізгі уақыты өтінімдерді іздеу, қатысушыларды салыстыру және залдың бос уақытын тексеру сияқты қайталанатын әрекеттерден босайды. Әкімші өтінімдерді бір тізімнен қарап, сүзгілеу арқылы қажетті жазбаны жылдам табады.",
        "Пайдаланушы үшін тиімділік жеке кабинет пен хабарлама жүйесі арқылы байқалады. Ол тіркелген іс-шараларын, залға берген өтінімдерін және мәртебе өзгерістерін бір жерде көреді. Бұл ақпараттың жоғалуын азайтып, қызмет сапасын жақсартады.",
        "Жүйе болашақта QR check-in, email/SMS хабарламалар, онлайн төлем, аналитикалық есептер, зал жүктемесін болжау және сыртқы күнтізбелермен синхрондау модульдерімен толықтырылуы мүмкін. Осы кеңейтулер енгізілген жағдайда платформа тек дипломдық жоба емес, нақты ұйымға арналған толыққанды басқару құралына айналады.",
        "Экономикалық бөлімнің қорытындысы бойынша, жобаны әзірлеуге кететін бастапқы шығын салыстырмалы түрде төмен, ал процестерді автоматтандырудан алынатын уақыт үнемі жоғары. Сондықтан веб-жүйе практикалық қолдануға жарамды және одан әрі дамытуға перспективалы деп бағаланады.",
    ]:
        add_body(doc, text)


def add_conclusion_and_backmatter(doc: Document) -> None:
    add_heading1(doc, "ҚОРЫТЫНДЫ")
    for text in [
        "Дипломдық жұмыста іс-шаралар мен конференцияларды тіркеу және басқаруға арналған ORDA Smart Event System веб-жүйесі қарастырылды. Жұмыстың негізгі мақсаты - іс-шараларды жариялау, қатысушыларды тіркеу, конференц-залдарды брондау, өтінімдерді өңдеу және әкімшілік бақылау процестерін бір ақпараттық ортада ұйымдастыру болды.",
        "Зерттеу барысында іс-шараларды басқару саласындағы цифрландырудың маңызы талданып, пайдаланушы, қатысушы және әкімші сценарийлері анықталды. Жүйенің функционалдық және функционалдық емес талаптары белгіленіп, клиент-сервер архитектурасы, REST API логикасы және MongoDB Atlas деректер қорының құрылымы сипатталды.",
        "Жоба frontend бөлігінде HTML, CSS және JavaScript технологияларын, backend бөлігінде Node.js және Express.js платформасын қолданады. Деректерді сақтау үшін MongoDB Atlas және Mongoose модельдері пайдаланылды. Авторизация JWT token арқылы ұйымдастырылып, әкімші функциялары рөлдік қолжетімділікпен қорғалды.",
        "Жүйеде іс-шаралар каталогы, конференц-залдар тізімі, календарь, пайдаланушы профилі, іс-шараға тіркелу, залды брондау, өтінім мәртебесін бақылау, хабарламалар, техникалық қолдау және әкімшілік панель сияқты негізгі модульдер іске асырылды. Көптілді интерфейс қазақ, орыс және ағылшын тілдерінде жұмыс істеуге мүмкіндік береді.",
        "Тестілеу нәтижесінде негізгі пайдаланушы сценарийлері орындалатыны анықталды: тіркеу, жүйеге кіру, іс-шараны таңдау, залға өтінім беру, admin панель арқылы мәртебе өзгерту және жеке кабинетте нәтижені көру. Бұл жүйенің дипломдық жоба деңгейінде толық жұмыс істейтінін көрсетеді.",
        "Еңбекті қорғау бөлімінде компьютермен жұмыс істеу қауіпсіздігі, жұмыс орнын ұйымдастыру және ақпараттық қауіпсіздік талаптары қарастырылды. Экономикалық бөлімде жобаның бастапқы шығындары мен практикалық тиімділігі есептелді.",
        "Қорытындылай келе, әзірленген веб-жүйе оқу орындары, конференц-орталықтар және ұйымдар үшін іс-шараларды басқару процесін жеңілдететін заманауи ақпараттық құрал болып табылады. Жоба болашақта онлайн төлем, QR check-in, email/SMS хабарламалар және кеңейтілген аналитика модульдерімен толықтырылуы мүмкін.",
    ]:
        add_body(doc, text)

    add_heading1(doc, "ПАЙДАЛАНЫЛҒАН ӘДЕБИЕТТЕР")
    refs = [
        "Node.js Documentation. URL: https://nodejs.org/docs",
        "Express.js Guide. URL: https://expressjs.com",
        "MongoDB Atlas Documentation. URL: https://www.mongodb.com/docs/atlas",
        "Mongoose Documentation. URL: https://mongoosejs.com/docs",
        "JSON Web Token Introduction. URL: https://jwt.io/introduction",
        "MDN Web Docs: HTML, CSS және JavaScript негіздері. URL: https://developer.mozilla.org",
        "REST API Design Best Practices. Microsoft Learn, 2025.",
        "OWASP Web Security Testing Guide. URL: https://owasp.org",
        "Қазақстан Республикасының ақпараттандыру туралы заңнамалық актілері.",
        "Ақпараттық жүйелерді жобалау негіздері: оқу құралы. Алматы, 2022.",
        "Веб-қосымшаларды әзірлеу технологиялары: оқу құралы. Астана, 2023.",
        "Бұлттық деректер қорларын қолдану тәжірибесі. Техникалық шолу, 2024.",
        "User Experience Design Principles for Web Applications. Nielsen Norman Group, 2024.",
        "Agile Web Development және бағдарламалық өнімді тестілеу әдістері. Оқу материалы, 2023.",
        "ORDA Smart Event System жобасының бастапқы коды және ішкі құжаттамасы, 2026.",
    ]
    for index, ref in enumerate(refs, start=1):
        add_reference(doc, f"{index}. {ref}")

    add_heading1(doc, "ҚОСЫМША А. БАҒДАРЛАМАНЫҢ НЕГІЗГІ КОДТАРЫ")
    for text in [
        "Бұл қосымшада ORDA Smart Event System жобасының негізгі бағдарламалық кодтары берілді. Код үзінділері жүйенің серверлік қосымшасын, деректер қоры модельдерін, маршруттарын, авторизация middleware қабатын, брондау логикасын және frontend helper модульдерін сипаттайды.",
        "Кодтар дипломдық жұмыстың жалпы стилін сақтай отырып құрылымдалды: алдымен файл атауы көрсетіледі, содан кейін жол нөмірлері бар негізгі бөлік беріледі. Мұндай рәсімдеу жобаның архитектурасын қорғау кезінде түсіндіруге және нақты модульдердің қызметін көрсетуге ыңғайлы.",
    ]:
        add_body(doc, text)
    code_files = [
        ("А.1 Сервердің негізгі іске қосылу бөлігі", "server.js", 80),
        ("А.2 Express қосымшасының конфигурациясы", "backend/src/app.js", 130),
        ("А.3 Орта айнымалылары және MongoDB қосылымы", "backend/src/config/env.js", 120),
        ("А.4 Деректер қорына қосылу модулі", "backend/src/config/db.js", 80),
        ("А.5 User моделі", "backend/src/models/User.js", 80),
        ("А.6 Event моделі", "backend/src/models/Event.js", 120),
        ("А.7 Hall моделі", "backend/src/models/Hall.js", 90),
        ("А.8 Booking моделі", "backend/src/models/Booking.js", 170),
        ("А.9 SupportMessage моделі", "backend/src/models/SupportMessage.js", 80),
        ("А.10 JWT авторизация middleware", "backend/src/middleware/auth.js", 100),
        ("А.11 Валидация және қате өңдеу middleware", "backend/src/middleware/validate.js", 90),
        ("А.12 API маршруттарын біріктіру", "backend/src/routes/index.js", 90),
        ("А.13 Авторизация маршруттары", "backend/src/routes/auth.routes.js", 80),
        ("А.14 Іс-шаралар маршруттары", "backend/src/routes/event.routes.js", 110),
        ("А.15 Конференц-зал маршруттары", "backend/src/routes/hall.routes.js", 90),
        ("А.16 Брондау маршруттары", "backend/src/routes/booking.routes.js", 140),
        ("А.17 Әкімші маршруттары", "backend/src/routes/admin.routes.js", 130),
        ("А.18 Техникалық қолдау маршруттары", "backend/src/routes/support.routes.js", 100),
        ("А.19 Авторизация controller бөлігі", "backend/src/controllers/auth.controller.js", 170),
        ("А.20 Іс-шара controller бөлігі", "backend/src/controllers/event.controller.js", 220),
        ("А.21 Зал controller бөлігі", "backend/src/controllers/hall.controller.js", 170),
        ("А.22 Брондау controller негізгі бөлігі", "backend/src/controllers/booking.controller.js", 360),
        ("А.23 Әкімші controller негізгі бөлігі", "backend/src/controllers/admin.controller.js", 260),
        ("А.24 Support controller бөлігі", "backend/src/controllers/support.controller.js", 120),
        ("А.25 Уақыт қақтығысын тексеру утилитасы", "backend/src/utils/availability.js", 260),
        ("А.26 Frontend HTTP helper", "docs/js/core/http.js", 100),
        ("А.27 Frontend session helper", "docs/js/core/session.js", 100),
        ("А.28 Көптілді интерфейс сөздігінің басы", "docs/js/core/i18n.js", 180),
        ("А.29 Frontend негізгі app логикасының басы", "docs/app.js", 320),
    ]
    for title, relative_path, max_lines in code_files:
        add_code_file(doc, title, relative_path, max_lines)


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_front_matter(doc)
    add_source_body(doc)
    add_labor_section(doc)
    add_economics_section(doc)
    add_conclusion_and_backmatter(doc)
    set_update_fields(doc)
    configure_document(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
