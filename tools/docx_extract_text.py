from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document


def iter_block_text(doc: Document):
    idx = 0
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            yield {
                "kind": "p",
                "index": idx,
                "style": paragraph.style.name if paragraph.style else "",
                "text": text,
            }
        idx += 1

    table_index = 0
    for table in doc.tables:
        for row_index, row in enumerate(table.rows):
            cells = [cell.text.strip().replace("\n", " | ") for cell in row.cells]
            text = " || ".join(cells).strip()
            if text:
                yield {
                    "kind": "table",
                    "index": table_index,
                    "row": row_index,
                    "style": "",
                    "text": text,
                }
        table_index += 1


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("--json", type=Path)
    parser.add_argument("--txt", type=Path)
    args = parser.parse_args()

    doc = Document(args.docx)
    blocks = list(iter_block_text(doc))

    if args.json:
        args.json.write_text(json.dumps(blocks, ensure_ascii=False, indent=2), encoding="utf-8")

    lines = []
    for block in blocks:
        if block["kind"] == "p":
            lines.append(f'{block["index"]:04d}\t{block["style"]}\t{block["text"]}')
        else:
            lines.append(f'T{block["index"]:02d}R{block["row"]:02d}\tTABLE\t{block["text"]}')

    output = "\n".join(lines)
    if args.txt:
        args.txt.write_text(output, encoding="utf-8")
    else:
        print(output)


if __name__ == "__main__":
    main()
