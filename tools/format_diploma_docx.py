from __future__ import annotations

import argparse
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from lxml import etree


FONT = "Times New Roman"
BODY_LINE_SPACING = 1.7
TABLE_FONT_SIZE = 14
TOP_BOTTOM_MARGIN_CM = 2.5
RIGHT_MARGIN_CM = 2
USABLE_WIDTH_CM = 16
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}


def set_run_font(run, size_pt: float | None = None, bold: bool | None = None) -> None:
    run.font.name = FONT
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(key), FONT)


def set_style_font(style, size_pt: float | None = None, bold: bool | None = None) -> None:
    style.font.name = FONT
    if size_pt is not None:
        style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(key), FONT)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def insert_paragraph_after(paragraph, text: str = "", style: str | None = None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def add_complex_toc_field(paragraph, levels: str = "1-3") -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_after = Pt(0)

    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f' TOC \\o "{levels}" \\h \\z \\u '

    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")

    placeholder = OxmlElement("w:t")
    placeholder.text = "Table of contents will update in Word"

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(placeholder)
    run._r.append(fld_end)


def add_page_field(paragraph) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "

    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")

    text = OxmlElement("w:t")
    text.text = "1"

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)
    set_run_font(run, 12)


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.widow_control = True


def style_document(doc: Document) -> None:
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(TOP_BOTTOM_MARGIN_CM)
        section.bottom_margin = Cm(TOP_BOTTOM_MARGIN_CM)
        section.left_margin = Cm(3)
        section.right_margin = Cm(RIGHT_MARGIN_CM)
        section.header_distance = Cm(1.25)
        section.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    set_style_font(normal, 14, False)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = BODY_LINE_SPACING
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    for name in ("Heading 1", "Heading 2", "Heading 3"):
        if name in doc.styles:
            set_style_font(doc.styles[name], 14, True)

    h1 = doc.styles["Heading 1"]
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.first_line_indent = Cm(0)
    h1.paragraph_format.line_spacing = BODY_LINE_SPACING
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.first_line_indent = Cm(0)
    h2.paragraph_format.line_spacing = BODY_LINE_SPACING
    h2.paragraph_format.space_before = Pt(6)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    if "FrontMatterHeading" not in doc.styles:
        front = doc.styles.add_style("FrontMatterHeading", WD_STYLE_TYPE.PARAGRAPH)
    else:
        front = doc.styles["FrontMatterHeading"]
    front.base_style = normal
    set_style_font(front, 14, True)
    front.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    front.paragraph_format.first_line_indent = Cm(0)
    front.paragraph_format.line_spacing = BODY_LINE_SPACING
    front.paragraph_format.space_before = Pt(12)
    front.paragraph_format.space_after = Pt(6)
    front.paragraph_format.keep_with_next = True

    if "CaptionCustom" in doc.styles:
        cap = doc.styles["CaptionCustom"]
    else:
        cap = doc.styles.add_style("CaptionCustom", WD_STYLE_TYPE.PARAGRAPH)
    cap.base_style = normal
    set_style_font(cap, 12, False)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.line_spacing = 1
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(6)
    cap.paragraph_format.keep_with_next = True


def normalize_front_matter_and_toc(doc: Document) -> None:
    h1_paragraphs = [p for p in doc.paragraphs if p.style and p.style.name == "Heading 1"]
    if len(h1_paragraphs) < 5:
        return

    for p in h1_paragraphs[:4]:
        p.style = "FrontMatterHeading"
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        set_keep_with_next(p)
        for run in p.runs:
            set_run_font(run, 14, True)

    contents_heading = h1_paragraphs[3]
    body_first_heading = h1_paragraphs[4]

    to_remove = []
    found_contents = False
    for p in doc.paragraphs:
        if p._p is contents_heading._p:
            found_contents = True
            continue
        if p._p is body_first_heading._p:
            break
        if found_contents:
            to_remove.append(p)

    for p in to_remove:
        remove_paragraph(p)

    toc = insert_paragraph_after(contents_heading)
    add_complex_toc_field(toc, "1-3")
    page_break = insert_paragraph_after(toc)
    page_break.paragraph_format.first_line_indent = Cm(0)
    page_break.paragraph_format.line_spacing = 1
    page_break.add_run().add_break(WD_BREAK.PAGE)


def normalize_paragraphs(doc: Document) -> None:
    h1_seen = 0
    body_started = False
    for p in doc.paragraphs:
        text = (p.text or "").strip()
        style_name = p.style.name if p.style is not None else ""

        if style_name == "Heading 1":
            h1_seen += 1
            if h1_seen >= 1:
                body_started = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = BODY_LINE_SPACING
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            set_keep_with_next(p)
            for run in p.runs:
                set_run_font(run, 14, True)
            continue

        if style_name == "Heading 2":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = BODY_LINE_SPACING
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            set_keep_with_next(p)
            for run in p.runs:
                set_run_font(run, 14, True)
            continue

        if style_name == "CaptionCustom" or text.lower().startswith(("figure", "table")):
            p.style = "CaptionCustom"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(6)
            set_keep_with_next(p)
            for run in p.runs:
                set_run_font(run, 12)
            continue

        if style_name == "List Bullet":
            p.paragraph_format.line_spacing = BODY_LINE_SPACING
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.first_line_indent = None
            p.paragraph_format.widow_control = True
            for run in p.runs:
                set_run_font(run, 14)
            continue

        # Preserve the title-page geometry; normalize the academic body after it.
        if body_started or h1_seen > 0:
            if text:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = BODY_LINE_SPACING
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.widow_control = True
            if style_name == "Normal":
                p.paragraph_format.first_line_indent = Cm(1.25)
            for run in p.runs:
                set_run_font(run, 14)


def normalize_tables_and_images(doc: Document) -> None:
    usable_width = Cm(USABLE_WIDTH_CM)
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        for row_idx, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.paragraph_format.line_spacing = 1
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    if len(table.columns) == 1:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif row_idx == 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        set_run_font(
                            run,
                            TABLE_FONT_SIZE,
                            bold=True if row_idx == 0 and len(table.rows) > 1 else None,
                        )

    for shape in doc.inline_shapes:
        if shape.width > usable_width:
            ratio = usable_width / shape.width
            shape.width = usable_width
            shape.height = int(shape.height * ratio)

    for p in doc.paragraphs:
        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p._p.xpath(".//w:drawing"):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p.paragraph_format.first_line_indent = Cm(0)


def normalize_footer(doc: Document) -> None:
    for section in doc.sections:
        section.different_first_page_header_footer = False
        footer = section.footer
        if not footer.paragraphs:
            p = footer.add_paragraph()
        else:
            p = footer.paragraphs[0]
        add_page_field(p)
        for extra in list(footer.paragraphs)[1:]:
            remove_paragraph(extra)


def prevent_blank_page_before_final_heading(doc: Document) -> None:
    headings = [p for p in doc.paragraphs if p.style is not None and p.style.name == "Heading 1"]
    if not headings:
        return
    final_heading = headings[-1]
    paragraphs = doc.paragraphs
    for idx, paragraph in enumerate(paragraphs):
        if paragraph._p is not final_heading._p or idx == 0:
            continue
        previous = paragraphs[idx - 1]
        page_breaks = previous._p.xpath('.//w:br[@w:type="page"]')
        if page_breaks and not (previous.text or "").strip():
            remove_paragraph(previous)
            final_heading.paragraph_format.page_break_before = False
        break


def set_update_fields_on_open(docx_path: Path) -> None:
    tmp_path = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/settings.xml":
                root = etree.fromstring(data)
                update = root.find("w:updateFields", namespaces=NS)
                if update is None:
                    update = etree.Element(f"{{{W_NS}}}updateFields")
                    root.insert(0, update)
                update.set(f"{{{W_NS}}}val", "true")
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    tmp_path.replace(docx_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    shutil.copyfile(args.input, args.output)
    doc = Document(str(args.output))
    style_document(doc)
    normalize_front_matter_and_toc(doc)
    normalize_paragraphs(doc)
    normalize_tables_and_images(doc)
    normalize_footer(doc)
    prevent_blank_page_before_final_heading(doc)
    doc.save(str(args.output))
    set_update_fields_on_open(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
